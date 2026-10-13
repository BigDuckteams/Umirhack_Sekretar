"""DOCX-шаблоны: проверка файла, поиск плейсхолдеров `{{ ключ }}`, заполнение и хранилище шаблонов.

Заполнение работает на уровне XML абзацев, поэтому видит плейсхолдеры в основном тексте, таблицах (в т.ч. вложенных),
надписях, колонтитулах всех разделов — и склеивает плейсхолдер, который Word разбил на несколько runs
(`{{ stu` + `dent_name }}`). Форматирование первого run плейсхолдера сохраняется. Исходный шаблон никогда не изменяется.
"""

from __future__ import annotations

import hashlib
import re
import secrets
import threading
import unicodedata
import uuid
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterator, Mapping, Optional

from docx import Document
from docx.oxml.ns import qn

PLACEHOLDER_RE = re.compile(r"\{\{\s*([^\W\d]\w*)\s*\}\}", re.UNICODE)

# ключи из ТЗ; шаблон может содержать и любые другие — они тоже заполняются, если данные переданы
STANDARD_KEYS = (
    "student_name", "passport_data", "current_university", "current_course", "current_specialty",
    "target_specialty", "education_form", "debts", "application_date",
)

MAX_TEMPLATE_BYTES = 10 * 1024 * 1024
MAX_UNPACKED_BYTES = 200 * 1024 * 1024
MAX_ENTRIES = 2000
_ID_RE = re.compile(r"[a-z0-9][a-z0-9_-]{1,63}")
_CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f￾￿]")
_XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
W_P, W_T, W_TC = qn("w:p"), qn("w:t"), qn("w:tc")


class TemplateError(ValueError):
    """Файл нельзя принять или обработать как шаблон (понятное пользователю сообщение)."""


class TemplateExists(TemplateError):
    """Шаблон с таким идентификатором уже есть, а перезапись не разрешена."""


class TemplateNotFound(TemplateError):
    pass


@dataclass
class PlaceholderInfo:
    key: str
    count: int = 0
    locations: list[str] = field(default_factory=list)  # body | table | header | footer
    known: bool = False  # входит в набор из ТЗ

    def to_dict(self) -> dict:
        return {"key": self.key, "count": self.count, "locations": self.locations, "known": self.known}


@dataclass
class TemplateInfo:
    template_id: str
    name: str
    version: int
    path: Path
    placeholders: list[PlaceholderInfo]
    sha256: str = ""
    size: int = 0

    @property
    def keys(self) -> list[str]:
        return [p.key for p in self.placeholders]

    @property
    def standard_keys(self) -> list[str]:
        return [p.key for p in self.placeholders if p.known]

    @property
    def custom_keys(self) -> list[str]:
        return [p.key for p in self.placeholders if not p.known]

    def to_dict(self) -> dict:
        return {"template_id": self.template_id, "name": self.name, "version": self.version, "sha256": self.sha256, "size": self.size,
                "placeholders": [p.to_dict() for p in self.placeholders], "standard_keys": self.standard_keys, "custom_keys": self.custom_keys}


@dataclass
class FillResult:
    path: Path
    filename: str
    filled: list[str]  # ключи, которые заменены на значения
    unfilled: list[str]  # плейсхолдеры, оставшиеся в документе (данных нет)
    unused: list[str]  # ключи данных, которых нет в шаблоне
    replaced: int = 0

    @property
    def complete(self) -> bool:
        return not self.unfilled

    def to_dict(self) -> dict:
        return {"filename": self.filename, "filled": self.filled, "unfilled": self.unfilled, "unused": self.unused, "complete": self.complete, "replaced": self.replaced}


# ====================================================================== проверка файла
def validate_docx(file_path: "str | Path", max_bytes: int = MAX_TEMPLATE_BYTES) -> Path:
    """Файл существует, это .docx, он безопасен для хранения и читается. Иначе TemplateError."""
    path = Path(file_path)
    if not path.exists() or not path.is_file():
        raise TemplateError(f"Файл не найден: {path.name or path}")
    if path.suffix.lower() != ".docx":
        raise TemplateError("Шаблон должен быть файлом .docx (документ Word без макросов)")
    size = path.stat().st_size
    if size == 0:
        raise TemplateError("Файл пустой")
    if size > max_bytes:
        raise TemplateError(f"Файл больше {max_bytes // (1024 * 1024)} МБ")
    if not zipfile.is_zipfile(path):
        raise TemplateError("Файл повреждён или не является документом Word (.docx)")
    try:
        with zipfile.ZipFile(path) as z:
            infos = z.infolist()
            names = {i.filename for i in infos}
            if len(infos) > MAX_ENTRIES or sum(i.file_size for i in infos) > MAX_UNPACKED_BYTES:
                raise TemplateError("Файл слишком большой после распаковки")
            for i in infos:
                if i.file_size > 1_000_000 and i.compress_size and i.file_size / i.compress_size > 500:
                    raise TemplateError("Файл выглядит как архив-бомба и не принят")
                if i.filename.startswith("/") or ".." in Path(i.filename).parts:
                    raise TemplateError("В файле некорректные пути")
            if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                raise TemplateError("Это не документ Word: нет word/document.xml")
            if any(n.lower().startswith(("word/vbaproject", "word/activex")) for n in names):
                raise TemplateError("В документе есть макросы или элементы ActiveX — такие файлы не принимаются")
            for n in names:
                if n.endswith(".rels"):
                    rels = z.read(n).decode("utf-8", "replace")
                    for rel in re.findall(r"<Relationship\s[^>]*>", rels):
                        if 'TargetMode="External"' in rel and "/hyperlink" not in rel:
                            raise TemplateError("Документ подключает внешний ресурс (шаблон, объект, рамку) по ссылке — такие файлы не принимаются")
    except zipfile.BadZipFile as exc:
        raise TemplateError("Файл повреждён: не удалось прочитать архив .docx") from exc
    try:
        Document(str(path))
    except Exception as exc:  # python-docx бросает разные исключения на битых файлах
        raise TemplateError(f"Не удалось прочитать документ Word: {exc}") from exc
    return path


# ====================================================================== обход XML
def _parts(doc) -> Iterator[tuple[str, Any]]:
    """(область, корневой XML-элемент): тело и все собственные колонтитулы разделов (общие не читаются дважды)."""
    yield "body", doc.element.body
    seen = set()
    for section in doc.sections:
        for kind, area in (("header", "header"), ("first_page_header", "header"), ("even_page_header", "header"),
                           ("footer", "footer"), ("first_page_footer", "footer"), ("even_page_footer", "footer")):
            hf = getattr(section, kind)
            if hf.is_linked_to_previous:  # своего содержимого нет — не создаём пустой колонтитул
                continue
            el = hf._element
            if id(el) not in seen:
                seen.add(id(el))
                yield area, el


def _own_text_nodes(p) -> list:
    """w:t этого абзаца; вложенные абзацы (надписи) обрабатываются как отдельные."""
    out, stack = [], list(reversed(list(p)))
    while stack:
        el = stack.pop()
        if el.tag == W_P:
            continue
        if el.tag == W_T:
            out.append(el)
            continue
        stack.extend(reversed(list(el)))
    return out


def _paragraphs(root) -> Iterator[Any]:
    yield from root.iter(W_P)


def _clean(value: Any) -> str:
    text = unicodedata.normalize("NFC", str(value))
    text = _CTRL.sub("", text)
    return re.sub(r"[\r\n\t]+", " ", text).strip()


def find_placeholders(source: "str | Path | Any") -> list[PlaceholderInfo]:
    """Плейсхолдеры документа: ключ, число вхождений, где встречается. Порядок — по первому появлению."""
    doc = source if hasattr(source, "element") else Document(str(source))
    found: dict[str, PlaceholderInfo] = {}
    for area, root in _parts(doc):
        for p in _paragraphs(root):
            nodes = _own_text_nodes(p)
            if not nodes:
                continue
            in_table = any(a.tag == W_TC for a in p.iterancestors())
            for m in PLACEHOLDER_RE.finditer("".join(n.text or "" for n in nodes)):
                info = found.setdefault(m.group(1), PlaceholderInfo(m.group(1), known=m.group(1) in STANDARD_KEYS))
                info.count += 1
                loc = "table" if in_table else area
                if loc not in info.locations:
                    info.locations.append(loc)
    return list(found.values())


# ====================================================================== заполнение
def _fill_paragraph(p, lookup: Mapping[str, str]) -> tuple[int, set[str]]:
    nodes = _own_text_nodes(p)
    if not nodes:
        return 0, set()
    texts = [n.text or "" for n in nodes]
    full = "".join(texts)
    matches = [m for m in PLACEHOLDER_RE.finditer(full) if m.group(1) in lookup]
    if not matches:
        return 0, set()
    starts, pos = [], 0
    for t in texts:
        starts.append(pos)
        pos += len(t)
    owner = [i for i, t in enumerate(texts) for _ in t]
    chars = [list(t) for t in texts]
    done: set[str] = set()
    for m in reversed(matches):
        s, e = m.start(), m.end()
        first, last = owner[s], owner[e - 1]
        ls, le = s - starts[first], e - starts[last]
        if first == last:
            chars[first][ls:le] = list(lookup[m.group(1)])
        else:
            chars[first][ls:] = list(lookup[m.group(1)])  # значение — в run, где начинался плейсхолдер (его форматирование)
            for k in range(first + 1, last):
                chars[k] = []
            chars[last][:le] = []
        done.add(m.group(1))
    for node, cs in zip(nodes, chars):
        text = "".join(cs)
        node.text = text
        if text != text.strip():
            node.set(_XML_SPACE, "preserve")
    return len(matches), done


def fill_docx(template_path: "str | Path", data: Mapping[str, Any], output_dir: "str | Path", *, blank_unfilled: bool = False,
              output_name: Optional[str] = None) -> FillResult:
    """Заполняет копию шаблона данными. Отсутствующие значения НЕ выдумываются: плейсхолдер остаётся в документе и попадает в `unfilled`
    (с `blank_unfilled=True` вместо него ставится «________»). Исходный шаблон не меняется."""
    template_path = Path(template_path)
    if not template_path.exists():
        raise TemplateNotFound(f"Шаблон не найден: {template_path.name}")
    try:
        doc = Document(str(template_path))
    except Exception as exc:
        raise TemplateError(f"Не удалось прочитать шаблон: {exc}") from exc
    lookup: dict[str, str] = {}
    for key, value in data.items():
        if value is None:
            continue
        text = _clean(value)
        if text:
            lookup[str(key)] = text
    in_template = {p.key for p in find_placeholders(doc)}
    blank = {k: "________" for k in in_template if k not in lookup} if blank_unfilled else {}
    effective = {**lookup, **blank}
    replaced, filled = 0, set()
    for _, root in _parts(doc):
        for p in _paragraphs(root):
            n, keys = _fill_paragraph(p, effective)
            replaced += n
            filled |= keys
    left = [p.key for p in find_placeholders(doc)]
    out_dir = Path(output_dir)
    try:
        out_dir.mkdir(parents=True, exist_ok=True)
        filename = output_name or f"{uuid.uuid4().hex}.docx"
        out_path = out_dir / filename
        doc.save(str(out_path))
    except OSError as exc:
        raise TemplateError(f"Не удалось сохранить документ: {exc}") from exc
    unfilled = list(dict.fromkeys([*left, *blank]))  # и оставшиеся метки, и заменённые прочерком: данных для них нет
    return FillResult(out_path, filename, sorted(filled - set(blank)), unfilled, sorted(set(lookup) - in_template), replaced)


# ====================================================================== хранилище шаблонов
def slugify(text: str) -> str:
    table = str.maketrans({"а": "a", "б": "b", "в": "v", "г": "g", "д": "d", "е": "e", "ё": "e", "ж": "zh", "з": "z", "и": "i", "й": "y", "к": "k", "л": "l",
                           "м": "m", "н": "n", "о": "o", "п": "p", "р": "r", "с": "s", "т": "t", "у": "u", "ф": "f", "х": "h", "ц": "c", "ч": "ch", "ш": "sh",
                           "щ": "sch", "ъ": "", "ы": "y", "ь": "", "э": "e", "ю": "yu", "я": "ya"})
    s = re.sub(r"[^a-z0-9]+", "-", text.lower().translate(table)).strip("-")
    return (s or "template")[:40]


class TemplateStore:
    """Каталог шаблонов `<templates_dir>/<id>/v<N>.docx`. С репозиторием (`storage.templates`) версии ведёт база, без него — по именам файлов."""

    def __init__(self, templates_dir: "str | Path", repo: Any = None):
        self.dir = Path(templates_dir)
        self.repo = repo
        self._lock = threading.Lock()  # две одновременные загрузки одного шаблона не должны писать в один файл версии

    # ---- пути
    def _resolve(self, rel: str) -> Path:
        path = (self.dir / rel).resolve()
        if self.dir.resolve() not in path.parents:
            raise TemplateError("Недопустимый путь к шаблону")
        return path

    def _file_versions(self, template_id: str) -> list[int]:
        d = self.dir / template_id
        if not d.is_dir():
            return []
        return sorted(int(m.group(1)) for f in d.glob("v*.docx") if (m := re.fullmatch(r"v(\d+)\.docx", f.name)))

    def exists(self, template_id: str) -> bool:
        if self.repo is not None:
            return self.repo.get(template_id) is not None
        return bool(self._file_versions(template_id))

    def path_of(self, template_id: str) -> Path:
        if self.repo is not None:
            row = self.repo.get(template_id)
            if row is None or not row.get("path"):
                raise TemplateNotFound("Шаблон не найден (возможно, его удалили)")
            path = self._resolve(row["path"])
        else:
            versions = self._file_versions(template_id)
            if not versions:
                raise TemplateNotFound("Шаблон не найден (возможно, его удалили)")
            path = self._resolve(f"{template_id}/v{versions[-1]}.docx")
        if not path.exists():
            raise TemplateNotFound("Файл шаблона отсутствует на диске — загрузите шаблон заново")
        return path

    # ---- добавление
    def add(self, file_path: "str | Path", *, template_id: Optional[str] = None, name: Optional[str] = None, overwrite: bool = False,
            actor: str = "system", note: str = "") -> TemplateInfo:
        src = validate_docx(file_path)  # долгая проверка — вне блокировки
        with self._lock:
            return self._add_checked(src, template_id, name, overwrite, actor, note)

    def _add_checked(self, src: Path, template_id: Optional[str], name: Optional[str], overwrite: bool, actor: str, note: str) -> TemplateInfo:
        placeholders = find_placeholders(src)
        display = (name or src.stem).strip()[:120] or "Шаблон"
        if template_id is not None:
            if not _ID_RE.fullmatch(template_id):
                raise TemplateError("Идентификатор шаблона: латинские строчные буквы, цифры, «-» и «_» (2–64 символа)")
            if self.exists(template_id) and not overwrite:
                raise TemplateExists(f"Шаблон «{template_id}» уже существует. Чтобы заменить его, укажите перезапись явно")
        else:
            template_id = f"{slugify(display)}-{secrets.token_hex(3)}"  # новый шаблон не может затереть чужой
            while self.exists(template_id):
                template_id = f"{slugify(display)}-{secrets.token_hex(3)}"
        data = src.read_bytes()
        sha = hashlib.sha256(data).hexdigest()
        if self.repo is not None:
            current = self.repo.get(template_id)
            version = (current["version"] if current else 0) + 1
        else:
            versions = self._file_versions(template_id)
            version = (versions[-1] if versions else 0) + 1
        rel = f"{template_id}/v{version}.docx"
        dest = self._resolve(rel)
        dest.parent.mkdir(parents=True, exist_ok=True)
        dest.write_bytes(data)
        if self.repo is not None:
            try:
                self.repo.add_version(template_id, display, src.name, rel, sha, len(data), [p.to_dict() for p in placeholders], actor, note)
            except Exception:
                dest.unlink(missing_ok=True)
                raise
        return TemplateInfo(template_id, display, version, dest, placeholders, sha, len(data))

    def delete(self, template_id: str) -> int:
        """Удаляет шаблон со всеми версиями и файлами. Возвращает число удалённых файлов."""
        if not self.exists(template_id):
            raise TemplateNotFound("Шаблон не найден")
        if self.repo is not None:
            rels = self.repo.delete(template_id)
        else:
            rels = [f"{template_id}/v{v}.docx" for v in self._file_versions(template_id)]
        n = 0
        for rel in rels:
            try:
                p = self._resolve(rel)
                if p.exists():
                    p.unlink()
                    n += 1
            except (OSError, TemplateError):
                continue
        d = self.dir / template_id
        try:
            d.rmdir()
        except OSError:
            pass
        return n


def convert_to_template(file_path: "str | Path", *, storage: Any = None, templates_dir: "str | Path | None" = None, template_id: Optional[str] = None,
                        name: Optional[str] = None, overwrite: bool = False, actor: str = "system") -> TemplateInfo:
    """Принимает эталонный DOCX (оформление, шапка, текст, подписи с плейсхолдерами `{{ ключ }}`), проверяет его и сохраняет как шаблон.

    * файл должен существовать, быть .docx и читаться; макросы и внешние шаблоны отклоняются;
    * шаблон кладётся в каталог шаблонов (`storage.templates_dir` из настроек либо `templates_dir`);
    * возвращается идентификатор шаблона и сведения о найденных плейсхолдерах;
    * чужой шаблон не перезаписывается: для замены нужны явные `template_id=...` и `overwrite=True`.
    """
    if storage is not None:
        store = TemplateStore(storage.templates_dir, storage.templates)
    else:
        if templates_dir is None:
            from ..config import load_settings

            templates_dir = load_settings().path("storage.templates_dir", "data/templates")
        store = TemplateStore(templates_dir)
    return store.add(file_path, template_id=template_id, name=name, overwrite=overwrite, actor=actor)
